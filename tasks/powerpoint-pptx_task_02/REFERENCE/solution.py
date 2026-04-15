"""
Reference solution for powerpoint-pptx_task_02.

Reads quarterly_sales.csv and produces:
  - sales_deck.pptx  (4 slides: title, table, chart, closing)
  - deck_report.json  (metadata about the deck)
"""
import csv
import json
from collections import defaultdict
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

workspace = Path(".")

# ---- Load CSV ---------------------------------------------------------------
rows = list(csv.DictReader((workspace / "quarterly_sales.csv").open()))
quarters = sorted(set(r["quarter"] for r in rows))
products = sorted(set(r["product"] for r in rows))

# Quarterly totals (for notes)
q_rev = defaultdict(float)
for r in rows:
    q_rev[r["quarter"]] += float(r["revenue"])

# Per-product per-quarter revenue dict
pq_rev = defaultdict(dict)
for r in rows:
    pq_rev[r["product"]][r["quarter"]] = float(r["revenue"])

total_revenue = sum(q_rev.values())

# ---- Build PPTX -------------------------------------------------------------
prs = Presentation()  # default 4:3 (9144000 × 6858000 EMU)
layout_map = {layout.name: layout for layout in prs.slide_layouts}

# Slide 0: Title Slide
slide0 = prs.slides.add_slide(layout_map["Title Slide"])
ph0 = {ph.placeholder_format.idx: ph for ph in slide0.placeholders}
ph0[0].text = "Annual Sales Review 2025"
ph0[1].text = "Quarterly performance by product line"

# Slide 1: Title Only + revenue table
slide1 = prs.slides.add_slide(layout_map["Title Only"])
ph1 = {ph.placeholder_format.idx: ph for ph in slide1.placeholders}
ph1[0].text = "Quarterly Revenue Summary"

# Table: 5 rows (header + 4 quarters), 4 cols (Quarter + 3 products)
tbl_left = 457200
tbl_top = 1143000
tbl_w = 8229600
tbl_h = 2743200
tbl_shape = slide1.shapes.add_table(
    5, 4, Emu(tbl_left), Emu(tbl_top), Emu(tbl_w), Emu(tbl_h)
)
tbl = tbl_shape.table
headers = ["Quarter"] + products
for c, h in enumerate(headers):
    tbl.cell(0, c).text = h
for r, q in enumerate(quarters, 1):
    tbl.cell(r, 0).text = q
    for c, p in enumerate(products, 1):
        rev = pq_rev[p].get(q, 0.0)
        tbl.cell(r, c).text = f"{rev:,.0f}"

# Speaker notes for slide 1
note1_parts = [f"{q}: ${q_rev[q]:,.0f}" for q in quarters]
slide1.notes_slide.notes_text_frame.text = (
    "Quarterly revenue grew steadily across all product lines. "
    + " | ".join(note1_parts)
)

# Slide 2: Title Only + clustered column chart
slide2 = prs.slides.add_slide(layout_map["Title Only"])
ph2 = {ph.placeholder_format.idx: ph for ph in slide2.placeholders}
ph2[0].text = "Revenue Trends by Product"

chart_data = ChartData()
chart_data.categories = quarters
for p in products:
    values = tuple(pq_rev[p].get(q, 0.0) for q in quarters)
    chart_data.add_series(p, values)

chart_shape = slide2.shapes.add_chart(
    XL_CHART_TYPE.COLUMN_CLUSTERED,
    Emu(457200), Emu(1143000), Emu(8229600), Emu(4572000),
    chart_data,
)

# Speaker notes for slide 2
top_product = max(products, key=lambda p: sum(pq_rev[p].values()))
slide2.notes_slide.notes_text_frame.text = (
    f"{top_product} drives the majority of annual revenue. "
    f"Total 2025 revenue: ${total_revenue:,.0f}."
)

# Slide 3: Closing (Title Slide layout)
slide3 = prs.slides.add_slide(layout_map["Title Slide"])
ph3 = {ph.placeholder_format.idx: ph for ph in slide3.placeholders}
ph3[0].text = "Thank You"
ph3[1].text = "questions@example.com | sales.example.com"

# ---- Save PPTX --------------------------------------------------------------
prs.save(workspace / "sales_deck.pptx")

# ---- Produce deck_report.json ------------------------------------------------
prs2 = Presentation(str(workspace / "sales_deck.pptx"))
slides_info = []
for i, slide in enumerate(prs2.slides):
    shape_types = [int(s.shape_type) for s in slide.shapes]
    has_chart = any(t == int(MSO_SHAPE_TYPE.CHART) for t in shape_types)
    has_table = any(t == int(MSO_SHAPE_TYPE.TABLE) for t in shape_types)
    notes_text = ""
    if slide.has_notes_slide:
        notes_text = slide.notes_slide.notes_text_frame.text.strip()
    chart_info = None
    if has_chart:
        for s in slide.shapes:
            if s.shape_type == MSO_SHAPE_TYPE.CHART:
                ch = s.chart
                chart_info = {
                    "chart_type": int(ch.chart_type),
                    "category_count": len(list(ch.plots[0].series[0].values)),
                    "series_count": len(list(ch.series)),
                    "series_names": [ser.name for ser in ch.series],
                }
                break
    slides_info.append({
        "slide_index": i,
        "layout_name": slide.slide_layout.name,
        "shape_count": len(slide.shapes),
        "shape_types": shape_types,
        "has_table": has_table,
        "has_chart": has_chart,
        "has_notes": bool(notes_text),
        "chart": chart_info,
    })

report = {
    "deck_title": "Annual Sales Review 2025",
    "slide_count": len(prs2.slides),
    "total_revenue": total_revenue,
    "quarters": quarters,
    "products": products,
    "quarterly_totals": {q: q_rev[q] for q in quarters},
    "slide_width_emu": prs2.slide_width,
    "slide_height_emu": prs2.slide_height,
    "slides": slides_info,
}
(workspace / "deck_report.json").write_text(json.dumps(report, indent=2))
print("Done — sales_deck.pptx and deck_report.json written")

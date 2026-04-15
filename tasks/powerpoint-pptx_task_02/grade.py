"""
Grade function for powerpoint-pptx_task_02.

Contract: returns a list of criterion records per grade-py-protocol.md.
Each record: {id, score in [0,1], weight (sum to 1.0), description, details?}.

Strategy: read sales_deck.pptx and deck_report.json from the workspace.
Compare against known expected values derived from quarterly_sales.csv fixture.

Archetypes tested:
  2 (common-default-wrong): chart series names and category labels must match CSV exactly.
  3 (multi-step coordination): title slide + table slide + chart slide + closing slide,
     all coherent with the same data source.
  5 (stateful invariant): chart category_count must equal quarter count (4),
     chart series_count must equal product count (3), report totals must match.
"""
from __future__ import annotations
import json
import math
import os
from pathlib import Path

# ---- Expected values derived from quarterly_sales.csv -----------------------

EXPECTED_SLIDE_COUNT = 4
EXPECTED_QUARTERS = ["Q1 2025", "Q2 2025", "Q3 2025", "Q4 2025"]
EXPECTED_PRODUCTS = ["Analytics Pro", "Data Bridge", "Insight API"]
EXPECTED_QUARTERLY_TOTALS = {
    "Q1 2025": 255500.0,
    "Q2 2025": 279600.0,
    "Q3 2025": 295000.0,
    "Q4 2025": 320500.0,
}
EXPECTED_TOTAL_REVENUE = 1150600.0

# Chart: COLUMN_CLUSTERED = 51 (XL_CHART_TYPE)
EXPECTED_CHART_TYPE = 51
EXPECTED_CHART_CATEGORIES = 4   # 4 quarters
EXPECTED_CHART_SERIES = 3       # 3 products

# Per-product per-quarter values (for deep chart verification)
EXPECTED_PQ = {
    "Analytics Pro": [142500.0, 158900.0, 167300.0, 181200.0],
    "Data Bridge":   [78200.0,  81500.0,  85100.0,  91000.0],
    "Insight API":   [34800.0,  39200.0,  42600.0,  48300.0],
}

# Table spot checks (row 0 is header, rows 1-4 are quarters)
EXPECTED_TABLE_HEADER = ["Quarter", "Analytics Pro", "Data Bridge", "Insight API"]
EXPECTED_TABLE_ROWS = {
    "Q1 2025": ["Q1 2025", "142,500", "78,200", "34,800"],
    "Q4 2025": ["Q4 2025", "181,200", "91,000", "48,300"],
}

MSO_PLACEHOLDER = 14
MSO_TABLE = 19
MSO_CHART = 3


# ---- Helpers -----------------------------------------------------------------

def _approx(a, b, tol=0.5):
    try:
        return math.isclose(float(a), float(b), abs_tol=tol)
    except (TypeError, ValueError):
        return False


def _load_pptx(workspace_path: str):
    path = Path(workspace_path) / "sales_deck.pptx"
    if not path.exists():
        return None, "sales_deck.pptx not found"
    try:
        from pptx import Presentation
        return Presentation(str(path)), None
    except Exception as e:
        return None, f"sales_deck.pptx failed to open: {e}"


def _load_report(workspace_path: str):
    path = Path(workspace_path) / "deck_report.json"
    if not path.exists():
        return None, "deck_report.json not found"
    try:
        return json.loads(path.read_text()), None
    except json.JSONDecodeError as e:
        return None, f"deck_report.json is not valid JSON: {e}"


def _slides(prs):
    return list(prs.slides) if prs else []


def _ph_map(slide):
    return {ph.placeholder_format.idx: ph for ph in slide.placeholders}


def _get_table_shape(slide):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for s in slide.shapes:
        if s.shape_type == MSO_SHAPE_TYPE.TABLE:
            return s
    return None


def _get_chart_shape(slide):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for s in slide.shapes:
        if s.shape_type == MSO_SHAPE_TYPE.CHART:
            return s
    return None


# ---- Criterion checks --------------------------------------------------------

def _check_pptx_exists(prs, report, slides):
    if prs is None:
        return 0.0, "sales_deck.pptx missing or corrupted"
    return 1.0, None


def _check_report_exists(prs, report, slides):
    if report is None:
        return 0.0, "deck_report.json missing or corrupted"
    return 1.0, None


def _check_slide_count(prs, report, slides):
    if prs is None:
        return 0.0, "sales_deck.pptx missing"
    n = len(slides)
    if n != EXPECTED_SLIDE_COUNT:
        return 0.0, f"expected {EXPECTED_SLIDE_COUNT} slides, got {n}"
    return 1.0, None


def _check_slide_layouts(prs, report, slides):
    """slide 0 and 3: 'Title Slide'; slide 1 and 2: 'Title Only'."""
    if prs is None or len(slides) < EXPECTED_SLIDE_COUNT:
        return 0.0, "sales_deck.pptx missing or incomplete"
    expected = ["Title Slide", "Title Only", "Title Only", "Title Slide"]
    errors = []
    for i, want in enumerate(expected):
        got = slides[i].slide_layout.name
        if got != want:
            errors.append(f"slide {i}: layout={got!r}, want {want!r}")
    if errors:
        return 0.0, "; ".join(errors)
    return 1.0, None


def _check_title_slide_text(prs, report, slides):
    """Slide 0 title must contain 'Annual Sales Review 2025', subtitle must be present."""
    if prs is None or len(slides) < 1:
        return 0.0, "sales_deck.pptx missing"
    ph = _ph_map(slides[0])
    title_ph = ph.get(0)
    sub_ph = ph.get(1)
    errors = []
    if title_ph is None:
        errors.append("slide 0: no title placeholder")
    elif "Annual Sales Review 2025" not in title_ph.text:
        errors.append(f"slide 0 title: expected 'Annual Sales Review 2025', got {title_ph.text!r}")
    if sub_ph is None:
        errors.append("slide 0: no subtitle placeholder")
    elif not sub_ph.text.strip():
        errors.append("slide 0 subtitle: empty")
    if errors:
        return 0.0, "; ".join(errors)
    return 1.0, None


def _check_table_exists(prs, report, slides):
    """Slide 1 must have a TABLE shape (type 19)."""
    if prs is None or len(slides) < 2:
        return 0.0, "sales_deck.pptx missing or too short"
    tbl_shape = _get_table_shape(slides[1])
    if tbl_shape is None:
        types = [int(s.shape_type) for s in slides[1].shapes]
        return 0.0, f"slide 1: no TABLE shape; shape_types={types}"
    return 1.0, None


def _check_table_header(prs, report, slides):
    """Table header row on slide 1 must exactly match ['Quarter', 'Analytics Pro', 'Data Bridge', 'Insight API']."""
    if prs is None or len(slides) < 2:
        return 0.0, "sales_deck.pptx missing"
    tbl_shape = _get_table_shape(slides[1])
    if tbl_shape is None:
        return 0.0, "slide 1: no table"
    tbl = tbl_shape.table
    got = [tbl.cell(0, c).text.strip() for c in range(min(4, len(tbl.columns)))]
    if got != EXPECTED_TABLE_HEADER:
        return 0.0, f"table header: expected {EXPECTED_TABLE_HEADER}, got {got}"
    return 1.0, None


def _check_table_row_count(prs, report, slides):
    """Table on slide 1 must have 5 rows (1 header + 4 quarter rows)."""
    if prs is None or len(slides) < 2:
        return 0.0, "sales_deck.pptx missing"
    tbl_shape = _get_table_shape(slides[1])
    if tbl_shape is None:
        return 0.0, "slide 1: no table"
    got = len(tbl_shape.table.rows)
    if got != 5:
        return 0.0, f"table: expected 5 rows (header + 4 quarters), got {got}"
    return 1.0, None


def _check_table_data_values(prs, report, slides):
    """Spot-check Q1 and Q4 rows in the table for revenue values matching the CSV."""
    if prs is None or len(slides) < 2:
        return 0.0, "sales_deck.pptx missing"
    tbl_shape = _get_table_shape(slides[1])
    if tbl_shape is None:
        return 0.0, "slide 1: no table"
    tbl = tbl_shape.table
    # Find rows by scanning the Quarter column
    errors = []
    q_row_map = {}
    for r in range(1, len(tbl.rows)):
        q = tbl.cell(r, 0).text.strip()
        q_row_map[q] = r
    for q, expected_row in EXPECTED_TABLE_ROWS.items():
        if q not in q_row_map:
            errors.append(f"table: row for {q!r} not found")
            continue
        r = q_row_map[q]
        for c, want_raw in enumerate(expected_row):
            got = tbl.cell(r, c).text.strip()
            # Normalize: strip commas for numeric columns
            got_norm = got.replace(",", "")
            want_norm = want_raw.replace(",", "")
            if got_norm != want_norm:
                errors.append(f"table cell({r},{c}): expected {want_raw!r}, got {got!r}")
    if errors:
        return 0.0, "; ".join(errors)
    return 1.0, None


def _check_chart_exists(prs, report, slides):
    """Slide 2 must have a CHART shape (type 3)."""
    if prs is None or len(slides) < 3:
        return 0.0, "sales_deck.pptx missing or too short"
    ch_shape = _get_chart_shape(slides[2])
    if ch_shape is None:
        types = [int(s.shape_type) for s in slides[2].shapes]
        return 0.0, f"slide 2: no CHART shape; shape_types={types}"
    return 1.0, None


def _check_chart_type(prs, report, slides):
    """Chart on slide 2 must be a clustered column chart (XL_CHART_TYPE.COLUMN_CLUSTERED = 51)."""
    if prs is None or len(slides) < 3:
        return 0.0, "sales_deck.pptx missing"
    ch_shape = _get_chart_shape(slides[2])
    if ch_shape is None:
        return 0.0, "slide 2: no chart"
    got = int(ch_shape.chart.chart_type)
    if got != EXPECTED_CHART_TYPE:
        return 0.0, f"chart type: expected COLUMN_CLUSTERED (51), got {got}"
    return 1.0, None


def _check_chart_category_count(prs, report, slides):
    """Chart must have exactly 4 categories (one per quarter) — a stateful invariant.
    Fewer categories means the agent dropped quarters; more means it added phantom data."""
    if prs is None or len(slides) < 3:
        return 0.0, "sales_deck.pptx missing"
    ch_shape = _get_chart_shape(slides[2])
    if ch_shape is None:
        return 0.0, "slide 2: no chart"
    try:
        series = list(ch_shape.chart.plots[0].series)
        if not series:
            return 0.0, "chart has no series"
        cat_count = len(list(series[0].values))
    except Exception as e:
        return 0.0, f"could not read chart categories: {e}"
    if cat_count != EXPECTED_CHART_CATEGORIES:
        return 0.0, f"chart category count: expected {EXPECTED_CHART_CATEGORIES}, got {cat_count}"
    return 1.0, None


def _check_chart_series_count(prs, report, slides):
    """Chart must have exactly 3 series (one per product) — a stateful invariant.
    If the agent collapsed products or added extra series, score fails."""
    if prs is None or len(slides) < 3:
        return 0.0, "sales_deck.pptx missing"
    ch_shape = _get_chart_shape(slides[2])
    if ch_shape is None:
        return 0.0, "slide 2: no chart"
    try:
        count = len(list(ch_shape.chart.series))
    except Exception as e:
        return 0.0, f"could not count series: {e}"
    if count != EXPECTED_CHART_SERIES:
        return 0.0, f"chart series count: expected {EXPECTED_CHART_SERIES} (one per product), got {count}"
    return 1.0, None


def _check_chart_series_names(prs, report, slides):
    """Chart series names must exactly match the three product names from the CSV.
    Models often paraphrase or truncate product names, which breaks the data contract."""
    if prs is None or len(slides) < 3:
        return 0.0, "sales_deck.pptx missing"
    ch_shape = _get_chart_shape(slides[2])
    if ch_shape is None:
        return 0.0, "slide 2: no chart"
    try:
        names = [s.name for s in ch_shape.chart.series]
    except Exception as e:
        return 0.0, f"could not read series names: {e}"
    missing = [p for p in EXPECTED_PRODUCTS if p not in names]
    if missing:
        return 0.0, f"chart missing series for products: {missing}; got names: {names}"
    return 1.0, None


def _check_chart_values(prs, report, slides):
    """Chart series values must match the CSV data for each product × quarter combination."""
    if prs is None or len(slides) < 3:
        return 0.0, "sales_deck.pptx missing"
    ch_shape = _get_chart_shape(slides[2])
    if ch_shape is None:
        return 0.0, "slide 2: no chart"
    errors = []
    try:
        series_map = {s.name: list(s.values) for s in ch_shape.chart.series}
    except Exception as e:
        return 0.0, f"could not read series values: {e}"
    for p, expected_vals in EXPECTED_PQ.items():
        if p not in series_map:
            errors.append(f"series {p!r} not found")
            continue
        got_vals = series_map[p]
        for i, (want, got) in enumerate(zip(expected_vals, got_vals)):
            if not _approx(want, got):
                errors.append(f"{p} Q{i+1}: expected {want}, got {got}")
    if errors:
        return 0.0, "; ".join(errors[:4])
    return 1.0, None


def _check_notes_on_slides_1_2(prs, report, slides):
    """Slides 1 and 2 must have non-empty speaker notes; slides 0 and 3 must be note-free."""
    if prs is None or len(slides) < EXPECTED_SLIDE_COUNT:
        return 0.0, "sales_deck.pptx missing or incomplete"
    expected_has_notes = {0: False, 1: True, 2: True, 3: False}
    errors = []
    for i, should_have in expected_has_notes.items():
        slide = slides[i]
        notes_text = ""
        if slide.has_notes_slide:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
        has = bool(notes_text)
        if should_have and not has:
            errors.append(f"slide {i}: expected notes, found none")
        elif not should_have and has:
            errors.append(f"slide {i}: expected no notes, found: {notes_text[:40]!r}")
    if errors:
        return 0.0, "; ".join(errors)
    return 1.0, None


def _check_closing_slide(prs, report, slides):
    """Slide 3 must use 'Title Slide' layout, title must contain 'Thank You', and subtitle must be non-empty."""
    if prs is None or len(slides) < EXPECTED_SLIDE_COUNT:
        return 0.0, "sales_deck.pptx missing"
    slide3 = slides[3]
    errors = []
    if slide3.slide_layout.name != "Title Slide":
        errors.append(f"closing slide layout: expected 'Title Slide', got {slide3.slide_layout.name!r}")
    ph = _ph_map(slide3)
    title_ph = ph.get(0)
    sub_ph = ph.get(1)
    if title_ph is None or "Thank You" not in title_ph.text:
        errors.append(f"closing slide title: expected 'Thank You', got {title_ph.text if title_ph else None!r}")
    if sub_ph is None or not sub_ph.text.strip():
        errors.append("closing slide subtitle: empty or missing")
    if errors:
        return 0.0, "; ".join(errors)
    return 1.0, None


def _check_report_slide_count(prs, report, slides):
    """deck_report.json slide_count must equal 4 — invariant between manifest and PPTX."""
    if report is None:
        return 0.0, "deck_report.json missing"
    got = report.get("slide_count")
    if got != EXPECTED_SLIDE_COUNT:
        return 0.0, f"report.slide_count: expected {EXPECTED_SLIDE_COUNT}, got {got}"
    return 1.0, None


def _check_report_total_revenue(prs, report, slides):
    """deck_report.json total_revenue must match the sum of all CSV revenue rows (1150600.0)."""
    if report is None:
        return 0.0, "deck_report.json missing"
    got = report.get("total_revenue")
    if not _approx(got, EXPECTED_TOTAL_REVENUE, tol=0.5):
        return 0.0, f"report.total_revenue: expected {EXPECTED_TOTAL_REVENUE}, got {got}"
    return 1.0, None


def _check_report_quarterly_totals(prs, report, slides):
    """deck_report.json quarterly_totals must report each quarter's summed revenue correctly."""
    if report is None:
        return 0.0, "deck_report.json missing"
    qt = report.get("quarterly_totals")
    if not isinstance(qt, dict):
        return 0.0, "report.quarterly_totals is missing or not an object"
    errors = []
    for q, want in EXPECTED_QUARTERLY_TOTALS.items():
        got = qt.get(q)
        if got is None:
            errors.append(f"quarterly_totals missing {q!r}")
        elif not _approx(got, want):
            errors.append(f"{q}: expected {want}, got {got}")
    if errors:
        return 0.0, "; ".join(errors)
    return 1.0, None


def _check_report_chart_metadata(prs, report, slides):
    """deck_report.json must report chart on slide 2 with correct category_count=4 and series_count=3."""
    if report is None:
        return 0.0, "deck_report.json missing"
    slides_data = report.get("slides", [])
    if len(slides_data) < 3:
        return 0.0, f"report.slides has only {len(slides_data)} entries"
    s2 = slides_data[2]
    ch = s2.get("chart")
    if not ch:
        return 0.0, "report.slides[2].chart is null or missing"
    errors = []
    if ch.get("category_count") != EXPECTED_CHART_CATEGORIES:
        errors.append(f"category_count: expected {EXPECTED_CHART_CATEGORIES}, got {ch.get('category_count')}")
    if ch.get("series_count") != EXPECTED_CHART_SERIES:
        errors.append(f"series_count: expected {EXPECTED_CHART_SERIES}, got {ch.get('series_count')}")
    if errors:
        return 0.0, "; ".join(errors)
    return 1.0, None


def _check_shape_type_consistency(prs, report, slides):
    """Shape type invariant: slide 1 has TABLE (19), slide 2 has CHART (3), others are PLACEHOLDER only.
    Text boxes (17) used instead of placeholders violate this invariant."""
    if prs is None or len(slides) < EXPECTED_SLIDE_COUNT:
        return 0.0, "sales_deck.pptx missing or incomplete"
    errors = []
    for i, slide in enumerate(slides):
        types = {int(s.shape_type) for s in slide.shapes}
        # Slides 0 and 3 should have only PLACEHOLDER shapes
        if i in (0, 3):
            unexpected = types - {MSO_PLACEHOLDER}
            if unexpected:
                errors.append(f"slide {i}: unexpected shape types {sorted(unexpected)} (expected only PLACEHOLDER)")
        # Slide 1 should have PLACEHOLDER + TABLE
        elif i == 1:
            if MSO_TABLE not in types:
                errors.append(f"slide 1: missing TABLE shape (19); got {sorted(types)}")
        # Slide 2 should have PLACEHOLDER + CHART
        elif i == 2:
            if MSO_CHART not in types:
                errors.append(f"slide 2: missing CHART shape (3); got {sorted(types)}")
    if errors:
        return 0.0, "; ".join(errors)
    return 1.0, None


# ---- Criterion registry -------------------------------------------------------

CRITERIA = [
    {
        "id": "pptx-exists",
        "weight": 0.03,
        "description": "sales_deck.pptx exists at the workspace root and opens without error.",
        "check": _check_pptx_exists,
    },
    {
        "id": "report-exists",
        "weight": 0.02,
        "description": "deck_report.json exists at the workspace root and is valid JSON.",
        "check": _check_report_exists,
    },
    {
        "id": "slide-count",
        "weight": 0.05,
        "description": "sales_deck.pptx contains exactly 4 slides: title, table, chart, closing.",
        "check": _check_slide_count,
    },
    {
        "id": "slide-layouts",
        "weight": 0.06,
        "description": "Slide layouts match the spec: slides 0 and 3 use 'Title Slide', slides 1 and 2 use 'Title Only'. Layout names must be matched by name, not by index.",
        "check": _check_slide_layouts,
    },
    {
        "id": "title-slide-text",
        "weight": 0.04,
        "description": "Slide 0 title contains 'Annual Sales Review 2025' and subtitle is non-empty, both in the correct placeholders (idx 0 and 1).",
        "check": _check_title_slide_text,
    },
    {
        "id": "table-exists",
        "weight": 0.04,
        "description": "Slide 1 contains a TABLE shape (MSO shape type 19). A text-only summary instead of a real table fails this check.",
        "check": _check_table_exists,
    },
    {
        "id": "table-header",
        "weight": 0.06,
        "description": "Table header row on slide 1 exactly matches ['Quarter', 'Analytics Pro', 'Data Bridge', 'Insight API'] — product names must match the CSV exactly, not be paraphrased.",
        "check": _check_table_header,
    },
    {
        "id": "table-row-count",
        "weight": 0.03,
        "description": "Table on slide 1 has 5 rows: 1 header row plus 4 quarter data rows matching the 4 distinct quarters in the CSV.",
        "check": _check_table_row_count,
    },
    {
        "id": "table-data-values",
        "weight": 0.06,
        "description": "Q1 2025 and Q4 2025 revenue values in the table match the CSV (e.g., Q4 Analytics Pro = 181,200). Values must come from the CSV, not be invented or averaged.",
        "check": _check_table_data_values,
    },
    {
        "id": "chart-exists",
        "weight": 0.04,
        "description": "Slide 2 contains a CHART shape (MSO shape type 3). A static image or table used as a chart substitute fails this check.",
        "check": _check_chart_exists,
    },
    {
        "id": "chart-type",
        "weight": 0.04,
        "description": "Chart on slide 2 is a clustered column chart (XL_CHART_TYPE.COLUMN_CLUSTERED = 51). Line charts, pie charts, or bar charts do not satisfy this requirement.",
        "check": _check_chart_type,
    },
    {
        "id": "chart-category-count",
        "weight": 0.08,
        "description": "Chart has exactly 4 categories (one per quarter). This is a stateful invariant: category_count must equal the number of distinct quarters in quarterly_sales.csv (4).",
        "check": _check_chart_category_count,
    },
    {
        "id": "chart-series-count",
        "weight": 0.08,
        "description": "Chart has exactly 3 series (one per product). This is a stateful invariant: series_count must equal the number of distinct products in the CSV (3). Collapsing products or adding aggregate series fails.",
        "check": _check_chart_series_count,
    },
    {
        "id": "chart-series-names",
        "weight": 0.07,
        "description": "Chart series names exactly match the three product names from the CSV: 'Analytics Pro', 'Data Bridge', 'Insight API'. Paraphrasing or truncating breaks downstream data contracts.",
        "check": _check_chart_series_names,
    },
    {
        "id": "chart-values",
        "weight": 0.07,
        "description": "Chart series values match the per-product per-quarter revenue from the CSV (e.g. Analytics Pro Q1=142500, Q4=181200). Values must be read from the CSV, not approximated.",
        "check": _check_chart_values,
    },
    {
        "id": "notes-slides-1-2",
        "weight": 0.04,
        "description": "Slides 1 and 2 have non-empty speaker notes; slides 0 and 3 have no notes. Notes placement must follow the spec and not be left empty on data slides.",
        "check": _check_notes_on_slides_1_2,
    },
    {
        "id": "closing-slide",
        "weight": 0.04,
        "description": "Slide 3 uses 'Title Slide' layout, its title contains 'Thank You', and its subtitle placeholder has contact or closing text — not left empty.",
        "check": _check_closing_slide,
    },
    {
        "id": "report-slide-count",
        "weight": 0.02,
        "description": "deck_report.json.slide_count equals 4, matching the actual PPTX slide count — a cross-artifact invariant.",
        "check": _check_report_slide_count,
    },
    {
        "id": "report-total-revenue",
        "weight": 0.04,
        "description": "deck_report.json.total_revenue equals 1150600.0, the sum of all revenue rows in quarterly_sales.csv. Confirms the agent computed the total from the data, not from the chart.",
        "check": _check_report_total_revenue,
    },
    {
        "id": "report-quarterly-totals",
        "weight": 0.04,
        "description": "deck_report.json.quarterly_totals reports the correct summed revenue per quarter (Q1=255500, Q2=279600, Q3=295000, Q4=320500), cross-validating both the table and the report.",
        "check": _check_report_quarterly_totals,
    },
    {
        "id": "report-chart-metadata",
        "weight": 0.03,
        "description": "deck_report.json.slides[2].chart reports category_count=4 and series_count=3 — a programmatic re-read of the PPTX proves the chart was written correctly, not just drawn.",
        "check": _check_report_chart_metadata,
    },
    {
        "id": "shape-type-consistency",
        "weight": 0.02,
        "description": "Shape type invariant across all slides: title slides have only PLACEHOLDER shapes (14), table slide has PLACEHOLDER + TABLE (19), chart slide has PLACEHOLDER + CHART (3). Text boxes (17) used in place of placeholders violate this.",
        "check": _check_shape_type_consistency,
    },
]

_w_sum = sum(c["weight"] for c in CRITERIA)
assert abs(_w_sum - 1.0) < 1e-3, f"CRITERIA weights sum to {_w_sum}, expected 1.0"


def grade(transcript, workspace_path):
    prs, _ = _load_pptx(workspace_path)
    report, _ = _load_report(workspace_path)
    slides = _slides(prs)

    records = []
    for spec in CRITERIA:
        score, details = spec["check"](prs, report, slides)
        record = {
            "id": spec["id"],
            "score": float(score),
            "weight": float(spec["weight"]),
            "description": spec["description"],
        }
        if details is not None and score < 1.0:
            record["details"] = details
        records.append(record)
    return records

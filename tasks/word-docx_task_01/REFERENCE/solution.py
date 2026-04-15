"""
Reference solution for word-docx_task_03.

Reads spec.json from cwd and writes report.docx using python-docx, exercising
named styles, a numbered list, a fixed-width table, a section break for the
appendix, and no tracked-change metadata.
"""
from __future__ import annotations

import json

from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Inches


def main() -> None:
    with open("spec.json") as f:
        spec = json.load(f)

    doc = Document()

    # --- Title -------------------------------------------------------------
    doc.add_paragraph(spec["title"], style="Title")

    # --- Overview ----------------------------------------------------------
    doc.add_paragraph("Overview", style="Heading 1")
    for para in spec["overview_paragraphs"]:
        doc.add_paragraph(para, style="Normal")

    # --- Action Items (real Word numbered list) ----------------------------
    doc.add_paragraph("Action Items", style="Heading 1")
    for item in spec["action_items"]:
        doc.add_paragraph(item, style="List Number")

    # --- Metrics table -----------------------------------------------------
    doc.add_paragraph("Metrics", style="Heading 1")
    table = doc.add_table(rows=1, cols=4)
    table.autofit = False
    col_width = Inches(1.5)

    # Set gridCol widths + every cell width so viewers don't re-auto-fit.
    for col in table.columns:
        col.width = col_width
    for cell in table.rows[0].cells:
        cell.width = col_width

    headers = spec["metrics_table"]["headers"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h

    for row_data in spec["metrics_table"]["rows"]:
        row = table.add_row()
        for i, v in enumerate(row_data):
            row.cells[i].text = v
            row.cells[i].width = col_width

    # --- Section break + Appendix -----------------------------------------
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph(spec["appendix_title"], style="Heading 1")
    for term, definition in spec["appendix_entries"]:
        doc.add_paragraph(term, style="Heading 2")
        doc.add_paragraph(definition, style="Normal")

    doc.save("report.docx")


if __name__ == "__main__":
    main()

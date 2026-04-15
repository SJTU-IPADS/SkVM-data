"""
Grade function for word-docx_task_03.

Returns a list of criterion records per docs/skvm/grade-py-protocol.md. Uses
python-docx to walk the generated report.docx and zipfile to inspect the raw
OOXML for tracked-change markers. No bun test, no JUnit.
"""
from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path


REQUIRED_HEADING1_TEXTS = ("overview", "action items", "metrics", "appendix")
MIN_HEADING2_COUNT = 2
MIN_LIST_NUMBER_COUNT = 3
TABLE_EXPECTED_ROWS = 4  # 1 header + 3 data
TABLE_EXPECTED_COLS = 4


def _safe_open_doc(workspace_path: str):
    path = Path(workspace_path) / "report.docx"
    if not path.exists():
        return None, "report.docx not found"
    try:
        from docx import Document
    except ImportError:
        return None, "python-docx not installed in grader environment"
    try:
        return Document(str(path)), None
    except Exception as e:
        return None, f"python-docx failed to open report.docx: {type(e).__name__}: {e}"


def _load_spec(workspace_path: str):
    for candidate in ("spec.json", "fixtures/spec.json"):
        p = Path(workspace_path) / candidate
        if p.exists():
            try:
                return json.loads(p.read_text())
            except json.JSONDecodeError:
                return None
    return None


def _iter_paragraphs_with_style(doc):
    return [(p.style.name if p.style else "", p.text.strip()) for p in doc.paragraphs]


# ---- per-criterion checks --------------------------------------------------

def _check_file_exists(ctx):
    if ctx["doc"] is None:
        return 0.0, ctx["doc_err"]
    return 1.0, None


def _check_valid_docx(ctx):
    if ctx["doc"] is None:
        return 0.0, ctx["doc_err"]
    # Also verify it's a legitimate DOCX zip with word/document.xml
    path = Path(ctx["cwd"]) / "report.docx"
    try:
        with zipfile.ZipFile(path) as zf:
            if "word/document.xml" not in zf.namelist():
                return 0.0, "report.docx is a zip but has no word/document.xml"
    except zipfile.BadZipFile:
        return 0.0, "report.docx is not a valid zip / docx package"
    return 1.0, None


def _check_title_content(ctx):
    if ctx["doc"] is None or ctx["spec"] is None:
        return 0.0, ctx["doc_err"] or "spec.json missing"
    title = ctx["spec"].get("title", "").strip()
    for _, text in ctx["paras"]:
        if title and title in text:
            return 1.0, None
    return 0.0, f"title text not found in any paragraph: {title!r}"


def _check_title_style(ctx):
    if ctx["doc"] is None or ctx["spec"] is None:
        return 0.0, ctx["doc_err"] or "spec.json missing"
    title = ctx["spec"].get("title", "").strip()
    for style, text in ctx["paras"]:
        if style == "Title" and title in text:
            return 1.0, None
    return 0.0, f"no paragraph with style 'Title' contains the title text"


def _check_heading1_sections(ctx):
    if ctx["doc"] is None:
        return 0.0, ctx["doc_err"]
    h1s = [t for s, t in ctx["paras"] if s == "Heading 1"]
    matched = []
    for needle in REQUIRED_HEADING1_TEXTS:
        if any(needle in t.lower() for t in h1s):
            matched.append(needle)
    missing = [n for n in REQUIRED_HEADING1_TEXTS if n not in matched]
    if missing:
        return 0.0, f"missing Heading 1 sections for: {missing}; found H1s: {h1s}"
    return 1.0, None


def _check_heading2_appendix(ctx):
    if ctx["doc"] is None or ctx["spec"] is None:
        return 0.0, ctx["doc_err"] or "spec.json missing"
    h2s = [t for s, t in ctx["paras"] if s == "Heading 2"]
    if len(h2s) < MIN_HEADING2_COUNT:
        return 0.0, f"expected at least {MIN_HEADING2_COUNT} Heading 2 paragraphs, got {len(h2s)}: {h2s}"
    expected_terms = [e[0] for e in ctx["spec"].get("appendix_entries", [])]
    missing = [term for term in expected_terms if not any(term == t or term in t for t in h2s)]
    if missing:
        return 0.0, f"missing Heading 2 text for appendix terms: {missing}; H2s: {h2s}"
    return 1.0, None


def _check_numbered_list_style(ctx):
    if ctx["doc"] is None or ctx["spec"] is None:
        return 0.0, ctx["doc_err"] or "spec.json missing"
    items = ctx["spec"].get("action_items", [])
    # Accept any of: "List Number", "List Number 2", "List Number 3", or a
    # style whose name starts with "List" (various built-ins / custom names).
    list_paras = [
        t for s, t in ctx["paras"]
        if s and (s == "List Number" or s.startswith("List ") or "Number" in s.split())
    ]
    if len(list_paras) < MIN_LIST_NUMBER_COUNT:
        return 0.0, (
            f"expected at least {MIN_LIST_NUMBER_COUNT} numbered-list paragraphs "
            f"(style 'List Number'), got {len(list_paras)}"
        )
    # Confirm each action item appears as one of the list paragraphs.
    for item in items:
        if not any(item[:40] in p for p in list_paras):
            return 0.0, (
                f"action item not found as numbered-list paragraph: {item[:60]!r}; "
                f"list paragraphs: {[p[:50] for p in list_paras]}"
            )
    return 1.0, None


def _check_table_exists(ctx):
    if ctx["doc"] is None:
        return 0.0, ctx["doc_err"]
    tables = ctx["doc"].tables
    if not tables:
        return 0.0, "no tables in report.docx"
    # Pick the first table with expected shape (there should be exactly one).
    for t in tables:
        if len(t.rows) == TABLE_EXPECTED_ROWS and len(t.columns) == TABLE_EXPECTED_COLS:
            return 1.0, None
    shapes = [(len(t.rows), len(t.columns)) for t in tables]
    return 0.0, f"no table has shape {TABLE_EXPECTED_ROWS}x{TABLE_EXPECTED_COLS}; found: {shapes}"


def _find_metrics_table(doc):
    for t in doc.tables:
        if len(t.rows) == TABLE_EXPECTED_ROWS and len(t.columns) == TABLE_EXPECTED_COLS:
            return t
    return None


def _check_table_explicit_widths(ctx):
    if ctx["doc"] is None:
        return 0.0, ctx["doc_err"]
    t = _find_metrics_table(ctx["doc"])
    if t is None:
        return 0.0, "no 4x4 metrics table found"
    missing_widths = []
    for i, col in enumerate(t.columns):
        w = col.width
        if w is None or getattr(w, "emu", w) == 0:
            missing_widths.append(i)
    if missing_widths:
        return 0.0, (
            f"columns {missing_widths} have no explicit width set (auto-fit); "
            f"SKILL.md rule: explicit table widths are safer than auto-fit"
        )
    return 1.0, None


def _check_table_content(ctx):
    if ctx["doc"] is None or ctx["spec"] is None:
        return 0.0, ctx["doc_err"] or "spec.json missing"
    t = _find_metrics_table(ctx["doc"])
    if t is None:
        return 0.0, "no 4x4 metrics table found"
    mt = ctx["spec"].get("metrics_table", {})
    expected_headers = mt.get("headers", [])
    expected_rows = mt.get("rows", [])

    header_cells = [c.text.strip() for c in t.rows[0].cells]
    for h in expected_headers:
        if h not in header_cells:
            return 0.0, f"table header missing: {h!r}; found: {header_cells}"

    body_text = "\n".join(
        "\t".join(cell.text.strip() for cell in row.cells)
        for row in t.rows[1:]
    )
    for row in expected_rows:
        for val in row:
            if val not in body_text:
                return 0.0, f"table cell value missing: {val!r}"
    return 1.0, None


def _check_two_sections(ctx):
    if ctx["doc"] is None:
        return 0.0, ctx["doc_err"]
    sections = list(ctx["doc"].sections)
    if len(sections) < 2:
        return 0.0, f"document has only {len(sections)} section(s), expected at least 2 (section break before appendix)"
    return 1.0, None


def _check_no_tracked_changes(ctx):
    if ctx["doc"] is None:
        return 0.0, ctx["doc_err"]
    path = Path(ctx["cwd"]) / "report.docx"
    try:
        with zipfile.ZipFile(path) as zf:
            xml = zf.read("word/document.xml").decode("utf-8", errors="replace")
    except Exception as e:
        return 0.0, f"could not read word/document.xml: {e}"
    # Look for actual tracked-change elements, not substring hits on names.
    for pattern in (r"<w:ins[\s>]", r"<w:del[\s>]", r"<w:commentRangeStart"):
        if re.search(pattern, xml):
            return 0.0, f"tracked-change element found matching {pattern!r}"
    return 1.0, None


def _check_overview_present(ctx):
    if ctx["doc"] is None or ctx["spec"] is None:
        return 0.0, ctx["doc_err"] or "spec.json missing"
    body_text = "\n".join(t for _, t in ctx["paras"])
    for para in ctx["spec"].get("overview_paragraphs", []):
        if para[:60] not in body_text:
            return 0.0, f"overview paragraph not found: {para[:80]!r}"
    return 1.0, None


def _check_action_items_present(ctx):
    if ctx["doc"] is None or ctx["spec"] is None:
        return 0.0, ctx["doc_err"] or "spec.json missing"
    body_text = "\n".join(t for _, t in ctx["paras"])
    for item in ctx["spec"].get("action_items", []):
        if item[:40] not in body_text:
            return 0.0, f"action item not found: {item[:80]!r}"
    return 1.0, None


CRITERIA = [
    {
        "id": "file-exists",
        "weight": 0.03,
        "description": "report.docx exists at the workspace root.",
        "check": _check_file_exists,
    },
    {
        "id": "valid-docx",
        "weight": 0.05,
        "description": "report.docx is a valid OOXML package — a ZIP containing word/document.xml that python-docx can open without errors.",
        "check": _check_valid_docx,
    },
    {
        "id": "title-content",
        "weight": 0.04,
        "description": "Title text from spec.json appears in the document body.",
        "check": _check_title_content,
    },
    {
        "id": "title-style",
        "weight": 0.08,
        "description": "The document's title paragraph uses the built-in 'Title' style — SKILL.md rule: prefer named styles over direct formatting.",
        "check": _check_title_style,
    },
    {
        "id": "heading1-sections",
        "weight": 0.10,
        "description": "All four top-level sections (Overview, Action Items, Metrics, Appendix) are rendered as 'Heading 1' paragraphs — a styles-layer check, not just visible text.",
        "check": _check_heading1_sections,
    },
    {
        "id": "heading2-appendix-entries",
        "weight": 0.06,
        "description": "Each appendix entry term (MTTR, Deploy) uses 'Heading 2' style — nested styles layer must be used deliberately, not just direct bold formatting.",
        "check": _check_heading2_appendix,
    },
    {
        "id": "numbered-list-style",
        "weight": 0.15,
        "description": "Action items are rendered using Word's numbered-list style ('List Number' or equivalent) — not plain paragraphs starting with '1.', and not pasted bullet glyphs. SKILL.md rule: bullets/numbering belong to numbering definitions, not to paragraph text.",
        "check": _check_numbered_list_style,
    },
    {
        "id": "table-exists",
        "weight": 0.05,
        "description": "Document contains a metrics table with exactly 4 columns and 4 rows (1 header row + 3 data rows).",
        "check": _check_table_exists,
    },
    {
        "id": "table-explicit-widths",
        "weight": 0.12,
        "description": "Every column in the metrics table has an explicit width set (not auto-fit). SKILL.md rule: 'explicit table widths are safer than auto-fit' for cross-viewer compatibility.",
        "check": _check_table_explicit_widths,
    },
    {
        "id": "table-content",
        "weight": 0.08,
        "description": "Metrics table cells contain all header values and all data row values from spec.metrics_table.",
        "check": _check_table_content,
    },
    {
        "id": "at-least-two-sections",
        "weight": 0.10,
        "description": "Document has at least 2 sections (sectPr elements) — the appendix lives in its own section, per the SKILL.md rule 'page layout lives in sections'.",
        "check": _check_two_sections,
    },
    {
        "id": "no-tracked-changes",
        "weight": 0.08,
        "description": "word/document.xml contains no tracked-change elements (w:ins, w:del, w:commentRangeStart). The generated file must be a clean delivery, not carry revision metadata.",
        "check": _check_no_tracked_changes,
    },
    {
        "id": "overview-present",
        "weight": 0.03,
        "description": "Both overview paragraphs from spec.json appear in the document body.",
        "check": _check_overview_present,
    },
    {
        "id": "action-items-present",
        "weight": 0.03,
        "description": "All three action-item strings from spec.json appear in the document body.",
        "check": _check_action_items_present,
    },
]

_w_sum = sum(c["weight"] for c in CRITERIA)
assert abs(_w_sum - 1.0) < 1e-3, f"CRITERIA weights sum to {_w_sum}, expected 1.0"


def grade(transcript, workspace_path):
    doc, doc_err = _safe_open_doc(workspace_path)
    spec = _load_spec(workspace_path)
    paras = _iter_paragraphs_with_style(doc) if doc is not None else []
    ctx = {
        "cwd": workspace_path,
        "doc": doc,
        "doc_err": doc_err,
        "spec": spec,
        "paras": paras,
    }
    records = []
    for spec_c in CRITERIA:
        score, details = spec_c["check"](ctx)
        rec = {
            "id": spec_c["id"],
            "score": float(score),
            "weight": float(spec_c["weight"]),
            "description": spec_c["description"],
        }
        if details is not None and score < 1.0:
            rec["details"] = details
        records.append(rec)
    return records

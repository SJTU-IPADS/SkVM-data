"""
Reference solution for excel-xlsx_task_02.

Reads sales.xlsx from cwd. Produces:
  - commissions.xlsx  with PayrollSummary sheet + Totals row (SUM formulas)
  - report.json  with quarterly_breakdown, by_salesperson, and totals

Key traps this solution handles correctly:
1. SaleDate: parses both Excel serial integers (epoch 1899-12-30) and ISO strings
2. Amount: parses both plain floats and '$X,XXX.XX' formatted strings
3. Commission: 8% per transaction for amount >= 1000, 5% otherwise (per-transaction)
4. Totals row: SUM formulas, not hardcoded values
5. Quarterly breakdown: correctly bucketed using parsed dates

Run: python3 REFERENCE/solution.py   (from workspace containing sales.xlsx)
"""
from __future__ import annotations

import json
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font


def parse_amount(v) -> float:
    """Parse amount: float/int or string like '$1,234.56' or '1234.56'."""
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("$", "").replace(",", "")
    return float(s)


def parse_date(v) -> date | None:
    """Parse SaleDate: Excel serial integer or ISO string 'YYYY-MM-DD'."""
    if isinstance(v, (int, float)):
        # Excel serial -> Python date (1900 epoch with leap-day bug)
        epoch = date(1899, 12, 30)
        return epoch + timedelta(days=int(v))
    if isinstance(v, str):
        try:
            return date.fromisoformat(v)
        except ValueError:
            return None
    return None


def compute_commission(amount: float) -> float:
    """8% if amount >= 1000, 5% otherwise. Applied per transaction."""
    rate = 0.08 if amount >= 1000.0 else 0.05
    return round(amount * rate, 2)


def quarter_label(d: date) -> str:
    """'Q1 2025', 'Q2 2025', 'Q3 2025', or 'Q4 2025'."""
    q = (d.month - 1) // 3 + 1
    return f"Q{q} {d.year}"


def main():
    wb_in = openpyxl.load_workbook("sales.xlsx", data_only=True)
    ws_in = wb_in["Transactions"]

    rows = []
    for row in range(2, ws_in.max_row + 1):
        txn_id = ws_in.cell(row=row, column=1).value
        if txn_id is None:
            continue
        salesperson = ws_in.cell(row=row, column=2).value
        sale_date = parse_date(ws_in.cell(row=row, column=3).value)
        amount = parse_amount(ws_in.cell(row=row, column=4).value)
        commission = compute_commission(amount)
        rows.append({
            "txn_id": txn_id,
            "salesperson": salesperson,
            "sale_date": sale_date,
            "amount": amount,
            "commission": commission,
        })

    # Per-salesperson aggregation
    person_stats: dict[str, dict] = defaultdict(
        lambda: {"sales_count": 0, "total_sales": 0.0, "commission": 0.0}
    )
    for r in rows:
        sp = r["salesperson"]
        person_stats[sp]["sales_count"] += 1
        person_stats[sp]["total_sales"] = round(person_stats[sp]["total_sales"] + r["amount"], 2)
        person_stats[sp]["commission"] = round(person_stats[sp]["commission"] + r["commission"], 2)

    sorted_persons = sorted(person_stats.keys())

    # ---- commissions.xlsx ----
    wb_out = Workbook()
    ws_out = wb_out.active
    ws_out.title = "PayrollSummary"

    headers = ["Salesperson", "SalesCount", "TotalSales", "Commission"]
    for col, h in enumerate(headers, 1):
        ws_out.cell(row=1, column=col, value=h).font = Font(bold=True)

    data_start = 2
    for ridx, sp in enumerate(sorted_persons, data_start):
        s = person_stats[sp]
        ws_out.cell(row=ridx, column=1, value=sp)
        ws_out.cell(row=ridx, column=2, value=s["sales_count"])
        ws_out.cell(row=ridx, column=3, value=round(s["total_sales"], 2))
        ws_out.cell(row=ridx, column=4, value=round(s["commission"], 2))

    # Totals row with SUM formulas
    total_row = data_start + len(sorted_persons)
    ws_out.cell(row=total_row, column=1, value="Total").font = Font(bold=True)
    ws_out.cell(row=total_row, column=2, value=f"=SUM(B{data_start}:B{total_row-1})")
    ws_out.cell(row=total_row, column=3, value=f"=SUM(C{data_start}:C{total_row-1})")
    ws_out.cell(row=total_row, column=4, value=f"=SUM(D{data_start}:D{total_row-1})")

    wb_out.save("commissions.xlsx")

    # ---- report.json ----
    quarter_stats: dict[str, dict] = defaultdict(
        lambda: {"transaction_count": 0, "sales": 0.0, "commission": 0.0}
    )
    for r in rows:
        if r["sale_date"] is None:
            continue
        ql = quarter_label(r["sale_date"])
        quarter_stats[ql]["transaction_count"] += 1
        quarter_stats[ql]["sales"] = round(quarter_stats[ql]["sales"] + r["amount"], 2)
        quarter_stats[ql]["commission"] = round(quarter_stats[ql]["commission"] + r["commission"], 2)

    total_sales = round(sum(s["total_sales"] for s in person_stats.values()), 2)
    total_commission = round(sum(s["commission"] for s in person_stats.values()), 2)

    report = {
        "total_transactions": len(rows),
        "total_sales": total_sales,
        "total_commission": total_commission,
        "quarterly_breakdown": {
            ql: {
                "transaction_count": quarter_stats[ql]["transaction_count"],
                "sales": round(quarter_stats[ql]["sales"], 2),
                "commission": round(quarter_stats[ql]["commission"], 2),
            }
            for ql in sorted(quarter_stats.keys())
        },
        "by_salesperson": {
            sp: {
                "sales_count": person_stats[sp]["sales_count"],
                "total_sales": round(person_stats[sp]["total_sales"], 2),
                "commission": round(person_stats[sp]["commission"], 2),
            }
            for sp in sorted_persons
        },
    }

    Path("report.json").write_text(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()

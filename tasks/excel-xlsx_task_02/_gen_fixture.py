"""
Fixture generator for excel-xlsx_task_02.

Run offline to regenerate sales.xlsx. The output is committed statically
and this script is kept for reproducibility only — it is NOT re-executed at
task evaluation time.

Usage:
    python3 _gen_fixture.py   # writes fixtures/sales.xlsx
"""
import random
from datetime import date, timedelta
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

random.seed(20260412)

SALESPERSONS = ["Alice Chen", "Bob Martinez", "Carol Singh", "David Lee", "Eva Turner"]


def date_to_excel_serial(d: date) -> int:
    """Convert Python date to Excel serial number (1900 epoch with leap-day bug)."""
    epoch = date(1899, 12, 30)
    return (d - epoch).days


rows = []
start_date = date(2025, 1, 1)
end_date = date(2025, 12, 31)
date_range = (end_date - start_date).days

for i in range(80):
    salesperson = random.choice(SALESPERSONS)
    amount_val = round(random.uniform(100.0, 5000.0), 2)

    # Amount: 70% float, 20% '$X,XXX.XX' string, 10% plain string
    r = random.random()
    if r < 0.70:
        amount = amount_val
    elif r < 0.90:
        amount = f"${amount_val:,.2f}"
    else:
        amount = str(amount_val)

    # SaleDate: 60% ISO string, 40% Excel serial integer
    sale_date = start_date + timedelta(days=random.randint(0, date_range))
    if random.random() < 0.60:
        date_val = sale_date.strftime("%Y-%m-%d")
    else:
        date_val = date_to_excel_serial(sale_date)

    rows.append({
        "TransactionID": f"TXN{i+1001:04d}",
        "Salesperson": salesperson,
        "SaleDate": date_val,
        "Amount": amount,
    })

random.shuffle(rows)

wb = Workbook()
ws = wb.active
ws.title = "Transactions"

headers = ["TransactionID", "Salesperson", "SaleDate", "Amount"]
for col, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=col, value=h)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="375623")

for ridx, row in enumerate(rows, 2):
    ws.cell(row=ridx, column=1, value=row["TransactionID"])
    ws.cell(row=ridx, column=2, value=row["Salesperson"])
    ws.cell(row=ridx, column=3, value=row["SaleDate"])
    ws.cell(row=ridx, column=4, value=row["Amount"])

ws.column_dimensions["A"].width = 14
ws.column_dimensions["B"].width = 16
ws.column_dimensions["C"].width = 14
ws.column_dimensions["D"].width = 14
ws.freeze_panes = "A2"

import os
out_path = os.path.join(os.path.dirname(__file__), "fixtures", "sales.xlsx")
wb.save(out_path)
print(f"Saved {out_path} with {len(rows)} rows")

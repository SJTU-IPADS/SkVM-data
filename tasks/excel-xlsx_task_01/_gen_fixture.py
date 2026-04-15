"""
Fixture generator for excel-xlsx_task_01.

Run offline to regenerate inventory.xlsx. The output is committed statically
and this script is kept for reproducibility only — it is NOT re-executed at
task evaluation time.

Usage:
    python3 _gen_fixture.py  # writes fixtures/inventory.xlsx
"""
import random
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from collections import defaultdict

random.seed(20260412)

CATEGORIES = ["Electronics", "Apparel", "Furniture", "Outdoors"]

rows = []
used_codes = set()

for i in range(70):
    # ItemCode: 5-char string always starting with "0"
    num = random.randint(1000, 9999)
    code = f"0{num}"
    while code in used_codes:
        num = random.randint(1000, 9999)
        code = f"0{num}"
    used_codes.add(code)

    cat = random.choice(CATEGORIES)
    unit_cost = round(random.uniform(5.0, 299.99), 2)
    qty = random.randint(1, 500)
    # First 12 rows are Discontinued; rest are Active
    status = "Discontinued" if i < 12 else "Active"
    rows.append({
        "ItemCode": code,
        "Category": cat,
        "UnitCost": unit_cost,
        "Quantity": qty,
        "Status": status,
    })

random.shuffle(rows)

wb = Workbook()
ws = wb.active
ws.title = "StockItems"

headers = ["ItemCode", "Category", "UnitCost", "Quantity", "Status"]
for col, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=col, value=h)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="2E75B6")

for ridx, row in enumerate(rows, 2):
    code_cell = ws.cell(row=ridx, column=1, value=row["ItemCode"])
    code_cell.number_format = "@"   # force text storage
    ws.cell(row=ridx, column=2, value=row["Category"])
    ws.cell(row=ridx, column=3, value=row["UnitCost"])
    ws.cell(row=ridx, column=4, value=row["Quantity"])
    ws.cell(row=ridx, column=5, value=row["Status"])

ws.freeze_panes = "A2"
ws.column_dimensions["A"].width = 12
ws.column_dimensions["B"].width = 14
ws.column_dimensions["C"].width = 12
ws.column_dimensions["D"].width = 10
ws.column_dimensions["E"].width = 14

import os
out_path = os.path.join(os.path.dirname(__file__), "fixtures", "inventory.xlsx")
wb.save(out_path)
print(f"Saved {out_path}")

# Verify expected values
active = [r for r in rows if r["Status"] == "Active"]
cat_stats = defaultdict(lambda: {"count": 0, "total_cost": 0.0})
for r in active:
    cat = r["Category"]
    inv_val = round(r["UnitCost"] * r["Quantity"], 2)
    cat_stats[cat]["count"] += 1
    cat_stats[cat]["total_cost"] = round(cat_stats[cat]["total_cost"] + inv_val, 2)

print(f"Active rows: {len(active)}")
for cat in sorted(cat_stats.keys()):
    print(f"  {cat}: count={cat_stats[cat]['count']}, total_cost={cat_stats[cat]['total_cost']:.2f}")

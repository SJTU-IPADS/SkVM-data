"""
Reference solution for excel-xlsx_task_01.

Reads inventory.xlsx from cwd. Produces:
  - report.xlsx  with two sheets: ActiveItems + CategoryTotals
  - inventory.json  with summary stats

Key traps this solution handles correctly:
1. ItemCode leading zeros preserved as text (zfill + number_format="@")
2. Only Active items in ActiveItems sheet (Discontinued excluded)
3. InventoryValue written as Excel formula =C{row}*D{row}
4. TotalInventoryCost and AvgUnitCost written as SUMIF/AVERAGEIF formulas
5. inventory.json computed values match (no formula: computes in Python)

Run: python3 REFERENCE/solution.py   (from workspace containing inventory.xlsx)
"""
from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font


def main():
    wb_in = openpyxl.load_workbook("inventory.xlsx", data_only=True)
    ws_in = wb_in["StockItems"]

    raw_rows = []
    for row in range(2, ws_in.max_row + 1):
        item_code = ws_in.cell(row=row, column=1).value
        if item_code is None:
            continue
        # Preserve leading zeros — openpyxl may return int if cell format was numeric
        item_code = str(item_code).zfill(5)
        category = ws_in.cell(row=row, column=2).value
        unit_cost = float(ws_in.cell(row=row, column=3).value)
        quantity = int(ws_in.cell(row=row, column=4).value)
        status = ws_in.cell(row=row, column=5).value
        raw_rows.append({
            "ItemCode": item_code,
            "Category": category,
            "UnitCost": unit_cost,
            "Quantity": quantity,
            "Status": status,
        })

    # Filter to Active only
    active_rows = [r for r in raw_rows if r["Status"] == "Active"]

    wb_out = Workbook()

    # ---- Sheet 1: ActiveItems ----
    ws_active = wb_out.active
    ws_active.title = "ActiveItems"

    active_headers = ["ItemCode", "Category", "UnitCost", "Quantity", "InventoryValue"]
    for col, h in enumerate(active_headers, 1):
        ws_active.cell(row=1, column=col, value=h).font = Font(bold=True)

    data_start = 2
    for ridx, row in enumerate(active_rows, data_start):
        code_cell = ws_active.cell(row=ridx, column=1, value=row["ItemCode"])
        code_cell.number_format = "@"  # store as text
        ws_active.cell(row=ridx, column=2, value=row["Category"])
        ws_active.cell(row=ridx, column=3, value=row["UnitCost"])
        ws_active.cell(row=ridx, column=4, value=row["Quantity"])
        # InventoryValue: Excel formula
        ws_active.cell(row=ridx, column=5, value=f"=C{ridx}*D{ridx}")

    # ---- Sheet 2: CategoryTotals ----
    ws_cat = wb_out.create_sheet("CategoryTotals")
    cat_headers = ["Category", "ItemCount", "TotalInventoryCost", "AvgUnitCost"]
    for col, h in enumerate(cat_headers, 1):
        ws_cat.cell(row=1, column=col, value=h).font = Font(bold=True)

    categories = sorted(set(r["Category"] for r in active_rows))
    cat_counts = defaultdict(int)
    for r in active_rows:
        cat_counts[r["Category"]] += 1

    for cidx, cat in enumerate(categories, 2):
        ws_cat.cell(row=cidx, column=1, value=cat)
        ws_cat.cell(row=cidx, column=2, value=cat_counts[cat])  # literal integer
        # TotalInventoryCost: SUMIF over ActiveItems InventoryValue (column E)
        ws_cat.cell(row=cidx, column=3, value=f"=SUMIF(ActiveItems!B:B,A{cidx},ActiveItems!E:E)")
        # AvgUnitCost: AVERAGEIF over ActiveItems UnitCost (column C)
        ws_cat.cell(row=cidx, column=4, value=f"=AVERAGEIF(ActiveItems!B:B,A{cidx},ActiveItems!C:C)")

    wb_out.save("report.xlsx")

    # ---- inventory.json ----
    cat_stats = defaultdict(lambda: {"item_count": 0, "total_cost": 0.0})
    total_cost = 0.0
    for r in active_rows:
        cat = r["Category"]
        inv_val = round(r["UnitCost"] * r["Quantity"], 2)
        cat_stats[cat]["item_count"] += 1
        cat_stats[cat]["total_cost"] = round(cat_stats[cat]["total_cost"] + inv_val, 2)
        total_cost += inv_val

    summary = {
        "total_active_items": len(active_rows),
        "total_inventory_cost": round(total_cost, 2),
        "by_category": {
            cat: {
                "item_count": cat_stats[cat]["item_count"],
                "total_cost": round(cat_stats[cat]["total_cost"], 2),
            }
            for cat in sorted(cat_stats.keys())
        },
    }

    Path("inventory.json").write_text(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()

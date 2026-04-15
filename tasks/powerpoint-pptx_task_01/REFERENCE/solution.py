"""
Reference solution for powerpoint-pptx_task_01.
Reads deck_spec.json and produces:
  - product_deck.pptx
  - deck_manifest.json
"""
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE

workspace = Path(".")
spec = json.loads((workspace / "deck_spec.json").read_text())

prs = Presentation()
# Keep default 4:3 slide dimensions (9144000 x 6858000 EMU)

# Map layout names to layout objects
layout_map = {layout.name: layout for layout in prs.slide_layouts}

slide_records = []

for slide_spec in spec["slides"]:
    layout_name = slide_spec["layout"]
    layout = layout_map[layout_name]
    slide = prs.slides.add_slide(layout)

    # Build placeholder index map
    ph_map = {ph.placeholder_format.idx: ph for ph in slide.placeholders}

    # Set title placeholder (idx=0 always)
    if "title" in slide_spec and 0 in ph_map:
        ph_map[0].text = slide_spec["title"]

    # Set subtitle placeholder (idx=1, Title Slide layout)
    if "subtitle" in slide_spec and 1 in ph_map:
        ph_map[1].text = slide_spec["subtitle"]

    # Set bullet content placeholder (idx=1, Title and Content layout)
    if "bullets" in slide_spec and 1 in ph_map:
        tf = ph_map[1].text_frame
        tf.clear()
        for i, bullet in enumerate(slide_spec["bullets"]):
            if i == 0:
                tf.paragraphs[0].text = bullet
            else:
                p = tf.add_paragraph()
                p.text = bullet
                p.level = 0

    # Add table if specified
    if "table" in slide_spec:
        tbl_spec = slide_spec["table"]
        pos = tbl_spec["position"]
        rows_data = tbl_spec["rows"]
        rows = len(rows_data)
        cols = len(rows_data[0])
        tbl_shape = slide.shapes.add_table(
            rows, cols,
            Emu(pos["left_emu"]),
            Emu(pos["top_emu"]),
            Emu(pos["width_emu"]),
            Emu(pos["height_emu"])
        )
        table = tbl_shape.table
        for r, row in enumerate(rows_data):
            for c, cell_text in enumerate(row):
                table.cell(r, c).text = cell_text

    # Set speaker notes
    if slide_spec.get("notes"):
        notes_slide = slide.notes_slide
        notes_slide.notes_text_frame.text = slide_spec["notes"]

    # Record manifest entry
    shape_types = [int(s.shape_type) for s in slide.shapes]

    slide_records.append({
        "slide_index": slide_spec["slide_index"],
        "layout_name": layout_name,
        "title": slide_spec.get("title", ""),
        "shape_count": len(slide.shapes),
        "shape_types": shape_types,
        "has_table": any(s == int(MSO_SHAPE_TYPE.TABLE) for s in shape_types),
        "has_notes": bool(slide_spec.get("notes", "").strip()),
    })

prs.save(workspace / "product_deck.pptx")

manifest = {
    "deck_title": spec["title"],
    "slide_count": len(prs.slides),
    "slides": slide_records,
    "slide_width_emu": prs.slide_width,
    "slide_height_emu": prs.slide_height,
}
(workspace / "deck_manifest.json").write_text(json.dumps(manifest, indent=2))
print("Done — product_deck.pptx and deck_manifest.json written")
